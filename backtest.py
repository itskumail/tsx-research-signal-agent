"""
backtest.py
-----------
Standalone historical simulator. Runs unattended in the background.

Walks 5 years of daily bars per watchlist name (free, via yfinance),
calls signals.score_symbol() at each bar, simulates ATR-based exits,
and writes:
  C:\\Trading\\backtest_report.html   — single-file HTML report
  C:\\Trading\\backtest_report.xlsx   — Excel mirror
  logs/backtest_<YYYYMMDD_HHMMSS>.log — full run log

Sends a Telegram completion ping when done.

Honest scope notes (READ BEFORE TRUSTING NUMBERS):
  • News sentiment, earnings blackout, tradability filters and
    time-of-day gates are NOT simulated (we don't have historical
    bid/ask, intraday bars, or news streams). The report flags this
    so you know the backtest is "rubric + ATR exits only".
  • VIX gate IS simulated using historical ^VIX bars.
  • Sector ETF gate IS simulated using historical sector ETF bars
    (best effort — some Canadian sector ETFs only date back ~10 yrs).
  • RS-vs-XIC, weekly-trend confirmation, and bearish-regime min_score
    bump ARE simulated.

Usage:
  python backtest.py                       # full watchlist, 5 yrs
  python backtest.py --years 3             # shorter window
  python backtest.py --symbols RY,TD,BMO   # subset
"""
from __future__ import annotations
import sys
import os
import json
import math
import time
import warnings
from datetime import datetime, timedelta
from pathlib import Path

warnings.filterwarnings("ignore")

import pandas as pd
import numpy as np

try:
    import yfinance as yf
except ImportError:
    print("ERROR: yfinance missing. pip install yfinance", file=sys.stderr); sys.exit(1)

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except ImportError:
    plt = None  # plots disabled, tables still work

from indicators import add_all
from signals   import score_symbol, TIMEFRAME_RULES
import filters

TRADING_DIR = Path(r"C:\Trading")
LOGS_DIR    = TRADING_DIR / "logs"
LOGS_DIR.mkdir(parents=True, exist_ok=True)

REPORT_HTML = TRADING_DIR / "backtest_report.html"
REPORT_XLSX = TRADING_DIR / "backtest_report.xlsx"
TS_TAG      = datetime.now().strftime("%Y%m%d_%H%M%S")
RUN_LOG     = LOGS_DIR / f"backtest_{TS_TAG}.log"


def log(msg: str):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"
    print(line, flush=True)
    with RUN_LOG.open("a", encoding="utf-8") as f:
        f.write(line + "\n")


# ---------------------------------------------------------------
# Data loaders
# ---------------------------------------------------------------
def load_watchlist() -> list[str]:
    out = []
    p = TRADING_DIR / "watchlist.txt"
    if not p.exists():
        return []
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        out.append(line)
    return out


def fetch_bars(ticker: str, years: int) -> pd.DataFrame | None:
    """Pull daily bars from yfinance.

    v5.4.2: uses yf_symbols.yf_ticker_candidates() so Class-A/B and TSXV
    names resolve correctly. Tries primary (.TO or .V per KNOWN_TSXV)
    then falls back to the other exchange suffix if the first returns
    nothing. Logs which candidate succeeded.
    """
    from yf_symbols import yf_ticker_candidates
    end   = datetime.now()
    start = end - timedelta(days=int(years * 365.25) + 30)

    last_err = None
    for yf_sym in yf_ticker_candidates(ticker, currency="CAD"):
        try:
            df = yf.download(yf_sym, start=start, end=end, progress=False,
                             auto_adjust=True, interval="1d")
            if df is None or df.empty or len(df) < 60:
                continue
            if isinstance(df.columns, pd.MultiIndex):
                df.columns = [c[0].lower() for c in df.columns]
            else:
                df.columns = [c.lower() for c in df.columns]
            df = df[["open", "high", "low", "close", "volume"]].copy()
            if yf_sym != f"{ticker}.TO":
                log(f"  resolved {ticker} -> {yf_sym}")
            return df
        except Exception as e:
            last_err = e
            continue
    if last_err is not None:
        log(f"  fetch error {ticker}: {last_err}")
    return None


def fetch_reference(symbol: str, years: int) -> pd.DataFrame | None:
    """Reference series (no .TO suffix) — XIC.TO, ^VIX, etc."""
    end   = datetime.now()
    start = end - timedelta(days=int(years * 365.25) + 30)
    try:
        df = yf.download(symbol, start=start, end=end, progress=False,
                         auto_adjust=True, interval="1d")
        if df is None or df.empty:
            return None
        if isinstance(df.columns, pd.MultiIndex):
            df.columns = [c[0].lower() for c in df.columns]
        else:
            df.columns = [c.lower() for c in df.columns]
        return df
    except Exception:
        return None


# ---------------------------------------------------------------
# Walk + simulate one symbol
# ---------------------------------------------------------------
def simulate(symbol: str, df: pd.DataFrame, xic_df: pd.DataFrame | None,
             vix_df: pd.DataFrame | None, sector_etfs: dict[str, pd.DataFrame],
             timeframe: str = "1-DAY", max_hold_bars: int = 60) -> list[dict]:
    """
    Walk bar by bar, compute signal at each, simulate next-bar-open
    entry and ATR-based exits. Returns list of trade dicts.
    """
    rules = TIMEFRAME_RULES[timeframe]
    min_history = rules["min_history_days"]

    # Pre-compute indicators on the full series
    df_full = add_all(df.copy())

    trades = []
    in_position = False
    entry_idx   = None
    entry_price = stop_loss = take_profit = 0.0
    bars_held   = 0

    # Map symbol → sector ETF (use Canadian sector ETFs from filters.SECTOR_MAP)
    sec = filters.sector_etf_for(symbol)
    sec_df = sector_etfs.get(sec) if sec else None

    for i in range(min_history, len(df_full) - 1):
        bar = df_full.iloc[i]
        date = df_full.index[i]
        next_open = df_full["open"].iloc[i + 1]

        if in_position:
            bars_held += 1
            high = df_full["high"].iloc[i]
            low  = df_full["low"].iloc[i]
            # Exit logic: hit TP, hit SL, or max-hold reached
            if low <= stop_loss:
                trades.append({"symbol": symbol, "entry_date": df_full.index[entry_idx],
                               "entry_price": entry_price, "exit_date": date,
                               "exit_price": stop_loss, "outcome": "SL",
                               "return_pct": (stop_loss - entry_price) / entry_price,
                               "bars_held": bars_held})
                in_position = False; bars_held = 0
                continue
            if high >= take_profit:
                trades.append({"symbol": symbol, "entry_date": df_full.index[entry_idx],
                               "entry_price": entry_price, "exit_date": date,
                               "exit_price": take_profit, "outcome": "TP",
                               "return_pct": (take_profit - entry_price) / entry_price,
                               "bars_held": bars_held})
                in_position = False; bars_held = 0
                continue
            if bars_held >= max_hold_bars:
                close_p = float(bar["close"])
                trades.append({"symbol": symbol, "entry_date": df_full.index[entry_idx],
                               "entry_price": entry_price, "exit_date": date,
                               "exit_price": close_p, "outcome": "TIMEOUT",
                               "return_pct": (close_p - entry_price) / entry_price,
                               "bars_held": bars_held})
                in_position = False; bars_held = 0
            continue

        # --- Try to generate a new signal at bar i ---
        df_slice = df_full.iloc[: i + 1]
        xic_slice = xic_df.iloc[: i + 1] if xic_df is not None else None

        # Weekly resample for multi-TF check
        weekly_slice = None
        if len(df_slice) > 60:
            try:
                weekly_slice = df_slice.resample("W").last().dropna()
            except Exception:
                pass

        # Live-equivalent regime + gates (best effort, daily-only)
        # VIX gate
        vix_level = None
        if vix_df is not None and date in vix_df.index:
            try: vix_level = float(vix_df.loc[date, "close"])
            except Exception: pass

        # Sector ETF % change today
        sec_pct = None
        if sec_df is not None and date in sec_df.index:
            try:
                idx = sec_df.index.get_loc(date)
                if idx > 0:
                    sec_pct = float(sec_df["close"].iloc[idx]) / float(sec_df["close"].iloc[idx-1]) - 1.0
            except Exception: pass

        sector_pct_dict = {sec: sec_pct} if sec and sec_pct is not None else None

        sig = score_symbol(
            df_slice, symbol=symbol, timeframe=timeframe, currency="CAD",
            xic_df=xic_slice, weekly_df=weekly_slice,
            vix_level=vix_level,
            sector_etf_pct_change=sector_pct_dict,
            news_events=None,             # no historical news
            snapshot_quote=None,          # no historical bid/ask
            wti_pct_change=None,          # could add per-day WTI return if needed
            use_atr_exits=True,
            enforce_time_window=False,    # daily bars
            skip_external_io=True,        # bypass live Yahoo earnings HTTP
        )

        if sig.sendable:
            in_position = True
            entry_idx   = i + 1
            entry_price = float(next_open)
            stop_loss   = sig.stop_loss
            take_profit = sig.take_profit
            bars_held   = 0

    return trades


# ---------------------------------------------------------------
# Aggregate stats
# ---------------------------------------------------------------
def aggregate(trades: list[dict]) -> dict:
    if not trades:
        return {
            "n": 0, "wins": 0, "losses": 0, "timeouts": 0,
            "win_rate": 0.0, "avg_win": 0.0, "avg_loss": 0.0, "expectancy": 0.0,
            "profit_factor": 0.0, "total_return": 0.0,
            "best_trade": 0.0, "worst_trade": 0.0,
            "avg_bars_held": 0.0,
        }
    wins   = [t for t in trades if t["return_pct"] > 0]
    losses = [t for t in trades if t["return_pct"] <= 0]
    timeouts = [t for t in trades if t["outcome"] == "TIMEOUT"]
    n = len(trades)
    avg_win  = np.mean([t["return_pct"] for t in wins])   if wins   else 0.0
    avg_loss = np.mean([t["return_pct"] for t in losses]) if losses else 0.0
    sum_win  = sum(t["return_pct"] for t in wins)
    sum_loss = abs(sum(t["return_pct"] for t in losses))
    pf = sum_win / sum_loss if sum_loss > 0 else float("inf") if sum_win > 0 else 0.0
    return {
        "n": n,
        "wins": len(wins), "losses": len(losses), "timeouts": len(timeouts),
        "win_rate": len(wins) / n if n else 0.0,
        "avg_win": float(avg_win), "avg_loss": float(avg_loss),
        "expectancy": (len(wins)/n)*float(avg_win) + (len(losses)/n)*float(avg_loss) if n else 0.0,
        "profit_factor": float(pf),
        "total_return": float(sum(t["return_pct"] for t in trades)),
        "best_trade":  float(max((t["return_pct"] for t in trades), default=0.0)),
        "worst_trade": float(min((t["return_pct"] for t in trades), default=0.0)),
        "avg_bars_held": float(np.mean([t["bars_held"] for t in trades])) if trades else 0.0,
    }


def equity_curve(trades: list[dict]) -> pd.Series:
    if not trades:
        return pd.Series(dtype=float)
    df = pd.DataFrame(trades).sort_values("exit_date")
    df["cum_return"] = (1 + df["return_pct"]).cumprod()
    df = df.set_index("exit_date")
    return df["cum_return"]


def drawdown(curve: pd.Series) -> float:
    if curve.empty:
        return 0.0
    running_max = curve.cummax()
    dd = (curve / running_max) - 1.0
    return float(dd.min())


# ---------------------------------------------------------------
# HTML report
# ---------------------------------------------------------------
def render_report(all_trades: list[dict], per_symbol_stats: dict, overall: dict,
                  curve: pd.Series, by_outcome: dict, by_sector: dict,
                  start_date: datetime, end_date: datetime, n_symbols: int) -> Path:
    rows_per = []
    for sym, st in sorted(per_symbol_stats.items(), key=lambda x: -x[1]["total_return"]):
        if st["n"] == 0: continue
        rows_per.append(
            f"<tr><td>{sym}</td><td class='right'>{st['n']}</td>"
            f"<td class='right'>{st['win_rate']*100:.1f}%</td>"
            f"<td class='right'>{st['avg_win']*100:+.2f}%</td>"
            f"<td class='right'>{st['avg_loss']*100:+.2f}%</td>"
            f"<td class='right'>{st['profit_factor']:.2f}</td>"
            f"<td class='right'>{st['total_return']*100:+.2f}%</td></tr>"
        )

    rows_recent = []
    for t in sorted(all_trades, key=lambda x: x["exit_date"], reverse=True)[:30]:
        rc = "green" if t["return_pct"] > 0 else "red"
        rows_recent.append(
            f"<tr><td>{str(t['entry_date'])[:10]}</td>"
            f"<td>{t['symbol']}</td>"
            f"<td>{t['outcome']}</td>"
            f"<td class='right'>${t['entry_price']:.2f}</td>"
            f"<td class='right'>${t['exit_price']:.2f}</td>"
            f"<td class='right {rc}'>{t['return_pct']*100:+.2f}%</td>"
            f"<td class='right'>{t['bars_held']}</td></tr>"
        )

    sector_rows = []
    for sec, st in sorted(by_sector.items(), key=lambda x: -x[1]["total_return"]):
        sector_rows.append(
            f"<tr><td>{sec}</td><td class='right'>{st['n']}</td>"
            f"<td class='right'>{st['win_rate']*100:.1f}%</td>"
            f"<td class='right'>{st['total_return']*100:+.2f}%</td></tr>"
        )

    # Equity curve PNG (best-effort)
    curve_img_html = ""
    if plt is not None and not curve.empty:
        fig, ax = plt.subplots(figsize=(10, 3.5))
        curve.plot(ax=ax, color="#58a6ff", lw=1.6)
        ax.set_title("Equity Curve (cumulative)")
        ax.grid(alpha=0.3)
        ax.set_ylabel("Multiple of initial")
        plt.tight_layout()
        png_path = TRADING_DIR / "backtest_equity.png"
        fig.savefig(png_path, dpi=110, facecolor="white")
        plt.close(fig)
        curve_img_html = f'<img src="backtest_equity.png" style="max-width:100%; border-radius:6px;">'

    days = (end_date - start_date).days or 1
    cagr = (curve.iloc[-1] ** (365 / days) - 1) if not curve.empty else 0.0
    max_dd = drawdown(curve)

    html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8">
<title>Backtest Report</title><style>
body {{ background:#0e1116; color:#e6edf3; margin:0; padding:20px;
       font-family:-apple-system,"Segoe UI",sans-serif; }}
h1 {{ margin:0; font-size:1.5rem; }}
h2 {{ font-size:1.05rem; margin-top:0; }}
.muted {{ color:#8b949e; font-size:0.85rem; }}
.panel {{ background:#161b22; border:1px solid #30363d; border-radius:8px;
         padding:16px; margin:14px 0; }}
.grid {{ display:grid; grid-template-columns:repeat(auto-fit, minmax(180px, 1fr));
        gap:10px; }}
.kv {{ background:#0d1117; padding:10px; border-radius:6px; border:1px solid #30363d; }}
.kv .k {{ color:#8b949e; font-size:0.78rem; }}
.kv .v {{ font-size:1.2rem; font-weight:600; }}
table {{ width:100%; border-collapse:collapse; font-size:0.88rem; }}
th, td {{ padding:7px 10px; text-align:left; border-bottom:1px solid #30363d; }}
th {{ color:#8b949e; font-weight:500; }}
.green {{ color:#3fb950; }} .red {{ color:#f85149; }}
.right {{ text-align:right; }}
.warn {{ background:rgba(210,153,34,0.1); border-left:3px solid #d29922;
        padding:10px; border-radius:4px; }}
</style></head><body>

<h1>📊 Backtest Report — Canadian Watchlist v4.2</h1>
<div class="muted">Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} ·
  {start_date.date()} → {end_date.date()} · {n_symbols} symbols</div>

<div class="warn" style="margin-top:14px;">
<b>Honest scope:</b> This backtest simulates the 7-check rubric, ATR-based exits,
relative strength vs XIC, weekly-trend confirmation, VIX gate (when historical
^VIX is available), and sector-ETF gate (when historical sector ETF data is
available). It does <b>NOT</b> simulate news sentiment, earnings blackout,
tradability filters (spread/liquidity), or time-of-day windows — historical
data for those isn't available in this environment. Treat numbers as
directional, not precise.
</div>

<section class="panel">
<h2>Headline Stats</h2>
<div class="grid">
  <div class="kv"><div class="k">Total trades</div><div class="v">{overall['n']}</div></div>
  <div class="kv"><div class="k">Win rate</div><div class="v">{overall['win_rate']*100:.1f}%</div></div>
  <div class="kv"><div class="k">Avg win</div><div class="v green">{overall['avg_win']*100:+.2f}%</div></div>
  <div class="kv"><div class="k">Avg loss</div><div class="v red">{overall['avg_loss']*100:+.2f}%</div></div>
  <div class="kv"><div class="k">Expectancy</div><div class="v">{overall['expectancy']*100:+.2f}%</div></div>
  <div class="kv"><div class="k">Profit factor</div><div class="v">{overall['profit_factor']:.2f}</div></div>
  <div class="kv"><div class="k">Total return (sum)</div><div class="v">{overall['total_return']*100:+.1f}%</div></div>
  <div class="kv"><div class="k">CAGR (compounded)</div><div class="v">{cagr*100:+.2f}%</div></div>
  <div class="kv"><div class="k">Max drawdown</div><div class="v red">{max_dd*100:.2f}%</div></div>
  <div class="kv"><div class="k">Avg bars held</div><div class="v">{overall['avg_bars_held']:.1f}</div></div>
  <div class="kv"><div class="k">Best trade</div><div class="v green">{overall['best_trade']*100:+.2f}%</div></div>
  <div class="kv"><div class="k">Worst trade</div><div class="v red">{overall['worst_trade']*100:+.2f}%</div></div>
</div>
</section>

<section class="panel">
<h2>Equity Curve</h2>
{curve_img_html}
</section>

<section class="panel">
<h2>By Outcome</h2>
<table><thead><tr><th>Outcome</th><th class="right">Count</th><th class="right">Avg Return</th></tr></thead>
<tbody>
""" + "".join(
        f'<tr><td>{k}</td><td class="right">{v["count"]}</td><td class="right">{v["avg"]*100:+.2f}%</td></tr>'
        for k, v in by_outcome.items()
    ) + """</tbody></table></section>

<section class="panel">
<h2>By Sector</h2>
<table><thead><tr><th>Sector ETF</th><th class="right">Trades</th>
<th class="right">Win Rate</th><th class="right">Total Return</th></tr></thead>
<tbody>""" + "".join(sector_rows) + """</tbody></table></section>

<section class="panel">
<h2>Per-Symbol Performance</h2>
<table><thead><tr><th>Symbol</th><th class="right">Trades</th>
<th class="right">Win Rate</th><th class="right">Avg Win</th>
<th class="right">Avg Loss</th><th class="right">Profit Factor</th>
<th class="right">Total Return</th></tr></thead>
<tbody>""" + "".join(rows_per) + """</tbody></table></section>

<section class="panel">
<h2>30 Most Recent Trades</h2>
<table><thead><tr><th>Entry</th><th>Symbol</th><th>Outcome</th>
<th class="right">Entry $</th><th class="right">Exit $</th>
<th class="right">Return</th><th class="right">Bars</th></tr></thead>
<tbody>""" + "".join(rows_recent) + """</tbody></table></section>

</body></html>"""

    REPORT_HTML.write_text(html, encoding="utf-8")
    return REPORT_HTML


# ---------------------------------------------------------------
# Excel report
# ---------------------------------------------------------------
def render_excel(all_trades: list[dict], per_symbol_stats: dict, overall: dict):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return None

    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    rows = [
        ("Total trades", overall["n"]),
        ("Win rate",     overall["win_rate"]),
        ("Avg win",      overall["avg_win"]),
        ("Avg loss",     overall["avg_loss"]),
        ("Expectancy",   overall["expectancy"]),
        ("Profit factor", overall["profit_factor"]),
        ("Total return", overall["total_return"]),
        ("Best trade",   overall["best_trade"]),
        ("Worst trade",  overall["worst_trade"]),
        ("Avg bars held", overall["avg_bars_held"]),
    ]
    ws["A1"] = "Backtest Report — Canadian Watchlist v4.2"
    ws["A1"].font = Font(bold=True, size=14)
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 16

    ws = wb.create_sheet("Per Symbol")
    headers = ["Symbol", "Trades", "Win Rate", "Avg Win", "Avg Loss", "Profit Factor", "Total Return"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, (sym, st) in enumerate(sorted(per_symbol_stats.items(), key=lambda x: -x[1]["total_return"]), 2):
        ws.cell(row=i, column=1, value=sym)
        ws.cell(row=i, column=2, value=st["n"])
        ws.cell(row=i, column=3, value=st["win_rate"])
        ws.cell(row=i, column=4, value=st["avg_win"])
        ws.cell(row=i, column=5, value=st["avg_loss"])
        ws.cell(row=i, column=6, value=st["profit_factor"])
        ws.cell(row=i, column=7, value=st["total_return"])

    ws = wb.create_sheet("Trades")
    headers = ["Entry Date", "Exit Date", "Symbol", "Outcome", "Entry $", "Exit $", "Return %", "Bars Held"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, t in enumerate(sorted(all_trades, key=lambda x: x["entry_date"]), 2):
        ws.cell(row=i, column=1, value=str(t["entry_date"])[:10])
        ws.cell(row=i, column=2, value=str(t["exit_date"])[:10])
        ws.cell(row=i, column=3, value=t["symbol"])
        ws.cell(row=i, column=4, value=t["outcome"])
        ws.cell(row=i, column=5, value=t["entry_price"])
        ws.cell(row=i, column=6, value=t["exit_price"])
        ws.cell(row=i, column=7, value=t["return_pct"])
        ws.cell(row=i, column=8, value=t["bars_held"])

    wb.save(REPORT_XLSX)
    return REPORT_XLSX


# ---------------------------------------------------------------
# Main
# ---------------------------------------------------------------
def main():
    args = sys.argv[1:]
    years = 5
    syms_override = None
    for a in args:
        if a.startswith("--years"):
            try: years = int(a.split("=")[1] if "=" in a else args[args.index(a)+1])
            except Exception: pass
        if a.startswith("--symbols"):
            try:
                val = a.split("=")[1] if "=" in a else args[args.index(a)+1]
                syms_override = [s.strip().upper() for s in val.split(",")]
            except Exception: pass

    log(f"=== Backtest start ({years}-year window) ===")

    watchlist = syms_override or load_watchlist()
    if not watchlist:
        log("No watchlist found. Exiting."); return

    log(f"Watchlist: {len(watchlist)} symbols")
    log("Pulling reference series (XIC.TO, ^VIX, sector ETFs)...")
    xic_df = fetch_reference("XIC.TO", years)
    vix_df = fetch_reference("^VIX", years)
    sector_etfs = {}
    for etf in ["XEG.TO", "XGD.TO", "XFN.TO", "XIT.TO", "XMA.TO", "XCM.TO"]:
        d = fetch_reference(etf, years)
        if d is not None: sector_etfs[etf.replace(".TO", "")] = d

    log(f"  XIC bars:    {len(xic_df) if xic_df is not None else 0}")
    log(f"  VIX bars:    {len(vix_df) if vix_df is not None else 0}")
    log(f"  Sector ETFs: {len(sector_etfs)} loaded")

    # Walk every symbol
    all_trades = []
    per_symbol_stats = {}
    start_date = datetime.now() - timedelta(days=years*365)
    end_date   = datetime.now()

    for i, sym in enumerate(watchlist, 1):
        log(f"[{i}/{len(watchlist)}] {sym} ...")
        df = fetch_bars(sym, years)
        if df is None or len(df) < 60:
            log(f"  no data — skip"); continue
        try:
            tr = simulate(sym, df, xic_df, vix_df, sector_etfs)
        except Exception as e:
            log(f"  simulate error: {e}"); continue

        per_symbol_stats[sym] = aggregate(tr)
        all_trades.extend(tr)
        log(f"  trades: {len(tr)} | win_rate: {per_symbol_stats[sym]['win_rate']*100:.0f}% | "
            f"total: {per_symbol_stats[sym]['total_return']*100:+.1f}%")

    log("Aggregating overall stats...")
    overall = aggregate(all_trades)
    curve   = equity_curve(all_trades)

    # By outcome
    by_outcome = {}
    for o in ["TP", "SL", "TIMEOUT"]:
        sub = [t for t in all_trades if t["outcome"] == o]
        by_outcome[o] = {
            "count": len(sub),
            "avg":   float(np.mean([t["return_pct"] for t in sub])) if sub else 0.0,
        }

    # By sector
    by_sector_raw = {}
    for t in all_trades:
        sec = filters.sector_etf_for(t["symbol"]) or "OTHER"
        by_sector_raw.setdefault(sec, []).append(t)
    by_sector = {sec: aggregate(ts) for sec, ts in by_sector_raw.items()}

    log("Rendering reports...")
    render_report(all_trades, per_symbol_stats, overall, curve, by_outcome, by_sector,
                  start_date, end_date, len(watchlist))
    render_excel(all_trades, per_symbol_stats, overall)

    log(f"=== Done. Reports written:")
    log(f"  {REPORT_HTML}")
    log(f"  {REPORT_XLSX}")

    # Telegram completion ping
    try:
        from telegram_notify import send
        days = (end_date - start_date).days or 1
        cagr = (curve.iloc[-1] ** (365 / days) - 1) if not curve.empty else 0.0
        max_dd = drawdown(curve)
        msg = (
            f"✅ <b>Backtest complete</b>\n"
            f"━━━━━━━━━━━━━━━━━━━━━━━\n"
            f"Window: {years} years · {len(watchlist)} symbols\n"
            f"Trades: <b>{overall['n']}</b>\n"
            f"Win rate: <b>{overall['win_rate']*100:.1f}%</b>\n"
            f"Profit factor: <b>{overall['profit_factor']:.2f}</b>\n"
            f"Total return (sum): <b>{overall['total_return']*100:+.1f}%</b>\n"
            f"CAGR: <b>{cagr*100:+.2f}%</b>\n"
            f"Max DD: <b>{max_dd*100:.2f}%</b>\n"
            f"\nReport: dashboard.html + backtest_report.html\n"
            f"<i>NB: rubric + ATR exits only. News/earnings/tradability not simulated — see report header.</i>"
        )
        send(msg)
    except Exception as e:
        log(f"Telegram ping failed: {e}")


if __name__ == "__main__":
    main()
