"""
backtest_sweep.py
-----------------
Parameter-sensitivity sweep. Runs the same 5-year historical simulation
under N parameter variants and writes a comparison report.

Variants tested (12 combinations):
  min_score      ∈ {5, 6, 7}
  ATR multiples  ∈ {(1.5, 2.5), (2.0, 3.0)}
  ADX mode       ∈ {"off", "gate>20"}

For each variant, walks every watchlist symbol bar-by-bar (same engine
as backtest.py), aggregates trades, and writes one row to the comparison
report. The "winner" is whichever variant gives the best Calmar ratio
(CAGR / |MaxDD|), which balances return vs. risk.

Output:
  C:\\Trading\\sweep_report.html
  C:\\Trading\\sweep_report.xlsx
  Telegram ping when done.

Runs in background, ~1-2 hours wall time.
"""
from __future__ import annotations
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path
import warnings

warnings.filterwarnings("ignore")

# UTF-8 stdio
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import pandas as pd
import numpy as np

import yfinance as yf
from indicators import add_all
from signals   import score_symbol, TIMEFRAME_RULES
import filters

TRADING_DIR = Path(r"C:\Trading")
LOGS_DIR    = TRADING_DIR / "logs"
LOGS_DIR.mkdir(parents=True, exist_ok=True)
TS_TAG      = datetime.now().strftime("%Y%m%d_%H%M%S")
RUN_LOG     = LOGS_DIR / f"sweep_{TS_TAG}.log"

REPORT_HTML = TRADING_DIR / "sweep_report.html"
REPORT_XLSX = TRADING_DIR / "sweep_report.xlsx"


def log(msg: str):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"
    print(line, flush=True)
    with RUN_LOG.open("a", encoding="utf-8") as f:
        f.write(line + "\n")


# ---------------------------------------------------------------
# Variants definition
# ---------------------------------------------------------------
VARIANTS = []
for min_s in (5, 6, 7):
    for atr_sl, atr_tp in ((1.5, 2.5), (2.0, 3.0)):
        for adx_mode in ("off", "gate"):
            VARIANTS.append({
                "name":    f"min{min_s}_atr{atr_sl}-{atr_tp}_adx{adx_mode}",
                "min_score": min_s,
                "atr_sl":    atr_sl,
                "atr_tp":    atr_tp,
                "adx_mode":  adx_mode,
            })


# ---------------------------------------------------------------
# Data loaders (reused from backtest.py shape)
# ---------------------------------------------------------------
def load_watchlist() -> list[str]:
    out = []
    p = TRADING_DIR / "watchlist.txt"
    if not p.exists(): return []
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"): continue
        out.append(line)
    return out


def fetch_bars(ticker: str, years: int) -> pd.DataFrame | None:
    # v5.4.2: route via yf_symbols for correct Class-A/B + TSXV handling
    from yf_symbols import yf_ticker_candidates
    end   = datetime.now()
    start = end - timedelta(days=int(years * 365.25) + 30)
    for yf_sym in yf_ticker_candidates(ticker, currency="CAD"):
        try:
            df = yf.download(yf_sym, start=start, end=end,
                             progress=False, auto_adjust=True, interval="1d")
            if df is None or df.empty or len(df) < 60: continue
            if isinstance(df.columns, pd.MultiIndex):
                df.columns = [c[0].lower() for c in df.columns]
            else:
                df.columns = [c.lower() for c in df.columns]
            return df[["open","high","low","close","volume"]].copy()
        except Exception:
            continue
    return None


def fetch_reference(symbol: str, years: int) -> pd.DataFrame | None:
    end   = datetime.now()
    start = end - timedelta(days=int(years * 365.25) + 30)
    try:
        df = yf.download(symbol, start=start, end=end, progress=False,
                         auto_adjust=True, interval="1d")
        if df is None or df.empty: return None
        if isinstance(df.columns, pd.MultiIndex):
            df.columns = [c[0].lower() for c in df.columns]
        else:
            df.columns = [c.lower() for c in df.columns]
        return df
    except Exception:
        return None


# ---------------------------------------------------------------
# Walk one symbol under one variant
# ---------------------------------------------------------------
def simulate_one(symbol: str, df_full: pd.DataFrame,
                 xic_df, vix_df, sector_etfs, variant: dict,
                 max_hold_bars: int = 60) -> list[dict]:
    """df_full has indicators already added (add_all). Returns trades."""
    rules       = TIMEFRAME_RULES["1-DAY"]
    min_history = rules["min_history_days"]

    sec    = filters.sector_etf_for(symbol)
    sec_df = sector_etfs.get(sec) if sec else None

    trades = []
    in_pos = False
    entry_idx = None; entry_price = sl = tp = 0.0; bars_held = 0

    for i in range(min_history, len(df_full) - 1):
        date = df_full.index[i]

        if in_pos:
            bars_held += 1
            high = df_full["high"].iloc[i]
            low  = df_full["low"].iloc[i]
            if low <= sl:
                trades.append({"symbol": symbol, "entry_date": df_full.index[entry_idx],
                               "entry_price": entry_price, "exit_date": date,
                               "exit_price": sl, "outcome": "SL",
                               "return_pct": (sl - entry_price) / entry_price,
                               "bars_held": bars_held})
                in_pos = False; bars_held = 0; continue
            if high >= tp:
                trades.append({"symbol": symbol, "entry_date": df_full.index[entry_idx],
                               "entry_price": entry_price, "exit_date": date,
                               "exit_price": tp, "outcome": "TP",
                               "return_pct": (tp - entry_price) / entry_price,
                               "bars_held": bars_held})
                in_pos = False; bars_held = 0; continue
            if bars_held >= max_hold_bars:
                close_p = float(df_full["close"].iloc[i])
                trades.append({"symbol": symbol, "entry_date": df_full.index[entry_idx],
                               "entry_price": entry_price, "exit_date": date,
                               "exit_price": close_p, "outcome": "TIMEOUT",
                               "return_pct": (close_p - entry_price) / entry_price,
                               "bars_held": bars_held})
                in_pos = False; bars_held = 0
            continue

        df_slice  = df_full.iloc[: i + 1]
        xic_slice = xic_df.iloc[: i + 1] if xic_df is not None else None
        weekly_slice = None
        if len(df_slice) > 60:
            try: weekly_slice = df_slice.resample("W").last().dropna()
            except Exception: pass

        # VIX gate
        vix_level = None
        if vix_df is not None and date in vix_df.index:
            try: vix_level = float(vix_df.loc[date, "close"])
            except Exception: pass

        # Sector ETF % change today
        sec_pct = None
        if sec_df is not None and date in sec_df.index:
            try:
                idx = sec_df.index.get_loc(date)
                if idx > 0:
                    sec_pct = float(sec_df["close"].iloc[idx]) / float(sec_df["close"].iloc[idx-1]) - 1.0
            except Exception: pass
        sector_pct_dict = {sec: sec_pct} if sec and sec_pct is not None else None

        # Override per-variant min_score by adjusting the rules dict mutably
        original_min = TIMEFRAME_RULES["1-DAY"]["min_score"]
        TIMEFRAME_RULES["1-DAY"]["min_score"] = variant["min_score"]
        try:
            sig = score_symbol(
                df_slice, symbol=symbol, timeframe="1-DAY", currency="CAD",
                xic_df=xic_slice, weekly_df=weekly_slice,
                vix_level=vix_level,
                sector_etf_pct_change=sector_pct_dict,
                news_events=None, snapshot_quote=None,
                wti_pct_change=None,
                use_atr_exits=True,
                enforce_time_window=False,
                adx_mode=variant["adx_mode"],
                adx_threshold=20.0,
                atr_sl_mult=variant["atr_sl"],
                atr_tp_mult=variant["atr_tp"],
                skip_external_io=True,    # bypass Yahoo earnings HTTP
            )
        finally:
            TIMEFRAME_RULES["1-DAY"]["min_score"] = original_min

        if sig.sendable:
            in_pos = True
            entry_idx = i + 1
            entry_price = float(df_full["open"].iloc[i + 1])
            sl = sig.stop_loss; tp = sig.take_profit
            bars_held = 0

    return trades


# ---------------------------------------------------------------
# Aggregate stats
# ---------------------------------------------------------------
def aggregate(trades: list[dict]) -> dict:
    if not trades:
        return {"n":0,"win_rate":0.0,"avg_win":0.0,"avg_loss":0.0,
                "profit_factor":0.0,"total_return":0.0,"cagr":0.0,
                "max_dd":0.0,"calmar":0.0,"avg_bars":0.0}
    wins = [t for t in trades if t["return_pct"] > 0]
    losses = [t for t in trades if t["return_pct"] <= 0]
    n = len(trades)
    avg_win  = float(np.mean([t["return_pct"] for t in wins])) if wins else 0.0
    avg_loss = float(np.mean([t["return_pct"] for t in losses])) if losses else 0.0
    sw = sum(t["return_pct"] for t in wins)
    sl_ = abs(sum(t["return_pct"] for t in losses))
    pf = sw / sl_ if sl_ > 0 else (float("inf") if sw > 0 else 0.0)

    df = pd.DataFrame(trades).sort_values("exit_date")
    df["cum"] = (1 + df["return_pct"]).cumprod()
    final = float(df["cum"].iloc[-1])
    days  = max(1, (pd.to_datetime(df["exit_date"].iloc[-1]) -
                    pd.to_datetime(df["entry_date"].iloc[0])).days)
    cagr  = final ** (365 / days) - 1
    running_max = df["cum"].cummax()
    max_dd = float(((df["cum"] / running_max) - 1).min())
    calmar = cagr / abs(max_dd) if max_dd < 0 else 0.0

    return {
        "n": n,
        "win_rate":      sum(1 for t in trades if t["return_pct"] > 0) / n,
        "avg_win":       avg_win,
        "avg_loss":      avg_loss,
        "profit_factor": float(pf),
        "total_return":  float(sum(t["return_pct"] for t in trades)),
        "cagr":          float(cagr),
        "max_dd":        max_dd,
        "calmar":        float(calmar),
        "avg_bars":      float(np.mean([t["bars_held"] for t in trades])),
    }


# ---------------------------------------------------------------
# Run one full variant across watchlist
# ---------------------------------------------------------------
def run_variant(variant: dict, dataset, xic_df, vix_df, sector_etfs) -> dict:
    log(f"--- variant: {variant['name']} ---")
    all_trades = []
    for sym, df_full in dataset.items():
        try:
            tr = simulate_one(sym, df_full, xic_df, vix_df, sector_etfs, variant)
        except Exception as e:
            log(f"  err {sym}: {e}"); continue
        all_trades.extend(tr)
    stats = aggregate(all_trades)
    log(f"  {variant['name']}: n={stats['n']} wr={stats['win_rate']*100:.1f}% "
        f"PF={stats['profit_factor']:.2f} CAGR={stats['cagr']*100:+.1f}% "
        f"DD={stats['max_dd']*100:.1f}% Calmar={stats['calmar']:.2f}")
    return {"variant": variant, "stats": stats, "trades": len(all_trades)}


# ---------------------------------------------------------------
# Reports
# ---------------------------------------------------------------
def render_html(results: list[dict], years: int):
    rows = sorted(results, key=lambda r: -r["stats"]["calmar"])
    rank = {id(r): i for i, r in enumerate(rows)}

    table = []
    for r in results:
        v = r["variant"]; s = r["stats"]
        is_winner = rank.get(id(r), 99) == 0
        cls = ' style="background:rgba(63,185,80,0.15);"' if is_winner else ''
        table.append(
            f"<tr{cls}><td>{v['name']}</td>"
            f"<td class='right'>{v['min_score']}</td>"
            f"<td class='right'>{v['atr_sl']}/{v['atr_tp']}</td>"
            f"<td>{v['adx_mode']}</td>"
            f"<td class='right'>{s['n']}</td>"
            f"<td class='right'>{s['win_rate']*100:.1f}%</td>"
            f"<td class='right'>{s['profit_factor']:.2f}</td>"
            f"<td class='right'>{s['cagr']*100:+.1f}%</td>"
            f"<td class='right red'>{s['max_dd']*100:.1f}%</td>"
            f"<td class='right'><b>{s['calmar']:.2f}</b></td></tr>"
        )

    html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8">
<title>Sweep Report</title><style>
body{{background:#0e1116;color:#e6edf3;margin:0;padding:20px;font-family:-apple-system,"Segoe UI",sans-serif}}
h1{{font-size:1.5rem;margin:0}}.muted{{color:#8b949e;font-size:0.85rem}}
.panel{{background:#161b22;border:1px solid #30363d;border-radius:8px;padding:16px;margin:14px 0}}
table{{width:100%;border-collapse:collapse;font-size:0.88rem}}
th,td{{padding:7px 10px;text-align:left;border-bottom:1px solid #30363d}}
th{{color:#8b949e;font-weight:500}}.right{{text-align:right}}.red{{color:#f85149}}
</style></head><body>
<h1>📊 Parameter Sweep — Canadian Watchlist v4.5</h1>
<div class="muted">Generated {datetime.now():%Y-%m-%d %H:%M} · {years}-year window · {len(results)} variants</div>
<div class="panel">
<p><b>Winner</b> (highest Calmar = best risk-adjusted return) is highlighted green.
Calmar = CAGR / |MaxDD|. Sort by it for the most balanced calibration.</p>
<table><thead><tr>
<th>Variant</th><th class="right">MinScore</th><th class="right">ATR(SL/TP)</th>
<th>ADX</th><th class="right">Trades</th><th class="right">WinRate</th>
<th class="right">PF</th><th class="right">CAGR</th>
<th class="right">MaxDD</th><th class="right">Calmar</th></tr></thead>
<tbody>{"".join(table)}</tbody></table></div></body></html>"""
    REPORT_HTML.write_text(html, encoding="utf-8")


def render_excel(results: list[dict]):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError: return

    wb = Workbook(); ws = wb.active; ws.title = "Variants"
    headers = ["Variant","MinScore","ATR_SL","ATR_TP","ADX","Trades",
               "WinRate","ProfitFactor","CAGR","MaxDD","Calmar","AvgBars"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, r in enumerate(results, 2):
        v = r["variant"]; s = r["stats"]
        ws.cell(row=i, column=1, value=v["name"])
        ws.cell(row=i, column=2, value=v["min_score"])
        ws.cell(row=i, column=3, value=v["atr_sl"])
        ws.cell(row=i, column=4, value=v["atr_tp"])
        ws.cell(row=i, column=5, value=v["adx_mode"])
        ws.cell(row=i, column=6, value=s["n"])
        ws.cell(row=i, column=7, value=s["win_rate"])
        ws.cell(row=i, column=8, value=s["profit_factor"])
        ws.cell(row=i, column=9, value=s["cagr"])
        ws.cell(row=i, column=10, value=s["max_dd"])
        ws.cell(row=i, column=11, value=s["calmar"])
        ws.cell(row=i, column=12, value=s["avg_bars"])
    wb.save(REPORT_XLSX)


# ---------------------------------------------------------------
# Main
# ---------------------------------------------------------------
def main():
    years = 5
    if "--years" in sys.argv:
        try: years = int(sys.argv[sys.argv.index("--years") + 1])
        except Exception: pass

    log(f"=== Sweep start: {len(VARIANTS)} variants × {years}-year window ===")

    watchlist = load_watchlist()
    log(f"Watchlist: {len(watchlist)} symbols")

    log("Pulling reference series...")
    xic_df = fetch_reference("XIC.TO", years)
    vix_df = fetch_reference("^VIX", years)
    sector_etfs = {}
    for etf in ["XEG.TO","XGD.TO","XFN.TO","XIT.TO","XMA.TO","XCM.TO"]:
        d = fetch_reference(etf, years)
        if d is not None: sector_etfs[etf.replace(".TO","")] = d

    # Pre-load + add indicators ONCE per symbol (huge speedup vs backtest.py)
    log("Pre-loading bars + indicators per symbol (one-time)...")
    dataset = {}
    for i, sym in enumerate(watchlist, 1):
        df = fetch_bars(sym, years)
        if df is None or len(df) < 60:
            continue
        try:
            df_full = add_all(df.copy())
            dataset[sym] = df_full
        except Exception as e:
            log(f"  add_all err {sym}: {e}")
        if i % 10 == 0:
            log(f"  loaded {i}/{len(watchlist)}")
    log(f"Dataset: {len(dataset)} symbols ready.")

    # Run all variants
    results = []
    for v in VARIANTS:
        results.append(run_variant(v, dataset, xic_df, vix_df, sector_etfs))

    # Sort by Calmar desc
    results.sort(key=lambda r: -r["stats"]["calmar"])

    log("Rendering reports...")
    render_html(results, years)
    render_excel(results)
    log(f"Done. Reports written:\n  {REPORT_HTML}\n  {REPORT_XLSX}")

    # Telegram summary
    try:
        from telegram_notify import send
        winner = results[0]
        wv = winner["variant"]; ws = winner["stats"]
        msg_lines = [
            "✅ <b>Parameter sweep complete</b>",
            "━━━━━━━━━━━━━━━━━━━━━━━",
            f"Tested {len(results)} variants over {years} years.",
            "",
            f"🥇 <b>Winner (best Calmar):</b>",
            f"  • min_score: <b>{wv['min_score']}/7</b>",
            f"  • ATR exits: <b>{wv['atr_sl']}× SL / {wv['atr_tp']}× TP</b>",
            f"  • ADX mode:  <b>{wv['adx_mode']}</b>",
            "",
            f"📊 <b>Numbers:</b>",
            f"  • Trades: {ws['n']}",
            f"  • Win rate: {ws['win_rate']*100:.1f}%",
            f"  • Profit factor: {ws['profit_factor']:.2f}",
            f"  • CAGR: {ws['cagr']*100:+.2f}%",
            f"  • Max DD: {ws['max_dd']*100:.2f}%",
            f"  • Calmar: <b>{ws['calmar']:.2f}</b>",
            "",
            f"🥈 <b>Runner-up:</b> {results[1]['variant']['name']} (Calmar {results[1]['stats']['calmar']:.2f})",
            f"🥉 <b>Bronze:</b> {results[2]['variant']['name']} (Calmar {results[2]['stats']['calmar']:.2f})",
            "",
            "Full report: sweep_report.html / .xlsx",
            "<i>Tomorrow we lock in the winner's params.</i>",
        ]
        send("\n".join(msg_lines))
    except Exception as e:
        log(f"telegram err: {e}")


if __name__ == "__main__":
    main()
